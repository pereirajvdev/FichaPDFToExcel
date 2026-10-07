import re
import sys
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURAÇÕES
# ============================================================

COLUNAS_VALORES = [
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
    "13º salário", "Rescisão", "Total",
]

MESES = COLUNAS_VALORES[:12]

PADRAO_DINHEIRO = r"(?<![\d.])-?(?:\d{1,3}(?:\.\d{3})+|\d+),\d{2}"

# Nome do arquivo usado como modelo
ARQUIVO_TEMPLATE = Path(
    r"data\AnexoVMacro.xlsm"
)
# Nome da aba do template
NOME_ABA = "Plan1"


# ============================================================
# PDF
# ============================================================

def extrair_texto_pdf(caminho_pdf: Path) -> str:
    """Extrai o texto de todas as páginas do PDF."""

    paginas = []

    with pdfplumber.open(caminho_pdf) as pdf:
        for pagina in pdf.pages:
            texto = pagina.extract_text() or ""
            paginas.append(texto)

    return "\n".join(paginas)


# ============================================================
# VALORES
# ============================================================

def normalizar_valor(valor: str) -> float:
    """Converte valor brasileiro, como 1.234,56, para float."""

    return float(
        valor.replace(".", "").replace(",", ".")
    )


# ============================================================
# DADOS DO SERVIDOR
# ============================================================

def extrair_dados_servidor(texto: str) -> dict:
    """Extrai os dados do servidor do cabeçalho da ficha."""

    dados = {}

    padrao_servidor = re.search(
        r"(?m)^(\d{6})\s+(.+?)\s+Secretaria:",
        texto
    )

    if padrao_servidor:
        dados["Matrícula"] = padrao_servidor.group(1)
        dados["Nome"] = padrao_servidor.group(2).strip()

    padrao = re.search(
        r"Ficha Financeira Resumo Geral do Ano de (\d{4})",
        texto
    )

    if padrao:
        dados["Ano"] = int(padrao.group(1))

    padrao = re.search(
        r"Secretaria:\s*(.*?)\s+Conta:",
        texto
    )

    if padrao:
        dados["Secretaria"] = padrao.group(1).strip()

    padrao = re.search(
        r"Conta:\s*(\S+)\s+Ano:",
        texto
    )

    if padrao:
        dados["Conta"] = padrao.group(1)

    padrao = re.search(
        r"Cargo:\s*(.*?)\s+CTPS:",
        texto
    )

    if padrao:
        dados["Cargo"] = padrao.group(1).strip()

    padrao = re.search(
        r"CTPS:\s*(\S+)\s+Pis/Pasep:",
        texto
    )

    if padrao:
        dados["CTPS"] = padrao.group(1)

    padrao = re.search(
        r"Pis/Pasep:\s*(\S+)\s+CPF:",
        texto
    )

    if padrao:
        dados["PIS/PASEP"] = padrao.group(1)

    padrao = re.search(
        r"CPF:\s*(\S+)\s+Dt\. Nasc\.:",
        texto
    )

    if padrao:
        dados["CPF"] = padrao.group(1)

    padrao = re.search(
        r"Dt\. Nasc\.\s*:\s*(\S+)\s+Dt\. Admissão:\s*(\S+)",
        texto
    )

    if padrao:
        dados["Data de nascimento"] = padrao.group(1)
        dados["Data de admissão"] = padrao.group(2)

    return dados


def extrair_dados_servidor_modelo_antigo(texto: str) -> dict:
    """
    Extrai os dados do servidor no modelo antigo de ficha financeira.
    """

    dados = {}

    # --------------------------------------------------------
    # MATRÍCULA + NOME
    # --------------------------------------------------------

    padrao = re.search(
        r"(?m)^Matricula:\s*(\d+)\s+(.+?)\s*$",
        texto,
        flags=re.IGNORECASE
    )

    if padrao:
        dados["Matrícula"] = padrao.group(1)
        dados["Nome"] = padrao.group(2).strip()

    # --------------------------------------------------------
    # CARGO
    # --------------------------------------------------------

    padrao = re.search(
        r"(?m)^Cargo:\s*(.+?)\s*$",
        texto,
        flags=re.IGNORECASE
    )

    if padrao:
        dados["Cargo"] = padrao.group(1).strip()

    # --------------------------------------------------------
    # ANO
    # --------------------------------------------------------

    padrao = re.search(
        r"Relatório\s+da\s+Ficha\s+Financeira\s+(\d{4})",
        texto,
        flags=re.IGNORECASE
    )

    if padrao:
        dados["Ano"] = int(padrao.group(1))

    return dados


# ============================================================
# RUBRICAS
# ============================================================

def encontrar_divisor(texto: str) -> int:
    """
    Encontra a posição de 'Total de Vencimentos'.

    Tudo que aparece antes do divisor será considerado
    vencimento e tudo depois será considerado desconto.
    """

    padrao = re.search(
        r"Total\s+de\s+Vencimentos",
        texto,
        flags=re.IGNORECASE
    )

    if not padrao:
        raise ValueError(
            "Não foi possível encontrar 'Total de Vencimentos' no PDF."
        )

    return padrao.start()


def extrair_rubricas(texto: str) -> tuple[list[dict], list[dict]]:
    """
    Extrai as rubricas e separa em Vencimentos e Descontos.
    """

    rubricas_vencimentos = []
    rubricas_descontos = []

    padrao_valor = re.compile(PADRAO_DINHEIRO)

    divisor = encontrar_divisor(texto)

    for linha in texto.splitlines():

        linha = linha.strip()

        if not re.match(r"^\d{5}\s+", linha):
            continue

        codigo = linha[:5]
        restante = linha[5:].strip()

        valores_encontrados = list(
            padrao_valor.finditer(restante)
        )

        if len(valores_encontrados) != 15:
            print(
                "Aviso: rubrica ignorada por quantidade "
                f"inesperada de valores: {linha}"
            )
            continue

        descricao = restante[
            :valores_encontrados[0].start()
        ].strip()

        valores = [
            normalizar_valor(m.group())
            for m in valores_encontrados
        ]

        rubrica = {
            "Código": codigo,
            "Descrição": descricao,
            **dict(zip(COLUNAS_VALORES, valores)),
        }

        posicao_linha = texto.find(linha)

        if posicao_linha < divisor:
            rubricas_vencimentos.append(rubrica)
        else:
            rubricas_descontos.append(rubrica)

    return rubricas_vencimentos, rubricas_descontos


def extrair_rubricas_modelo_antigo(texto: str) -> tuple[list[dict], list[dict]]:
    """
    Extrai rubricas do modelo antigo de ficha financeira.

    No modelo antigo:
    - cada rubrica aparece dentro de um mês;
    - existe apenas um valor por rubrica;
    - os valores são distribuídos nas colunas dos meses;
    - vencimentos e descontos são identificados pelo código.
    """

    rubricas_vencimentos = {}
    rubricas_descontos = {}

    mes_atual = None

    mapa_meses = {
        "janeiro": "Janeiro",
        "fevereiro": "Fevereiro",
        "março": "Março",
        "abril": "Abril",
        "maio": "Maio",
        "junho": "Junho",
        "julho": "Julho",
        "agosto": "Agosto",
        "setembro": "Setembro",
        "outubro": "Outubro",
        "novembro": "Novembro",
        "dezembro": "Dezembro",
    }

    # Códigos que aparecem como descontos no modelo antigo.
    codigos_desconto = {
        "403",
        "413",
        "500",
        "798",
    }

    padrao_rubrica = re.compile(
        r"^(\d{3})\s+(.+?)\s+"
        r"(-?(?:\d{1,3}(?:\.\d{3})+|\d+),\d{2})$"
    )

    for linha in texto.splitlines():

        linha = linha.strip()

        # --------------------------------------------------------
        # IDENTIFICAR MÊS
        # --------------------------------------------------------

        mes_normalizado = linha.lower()

        if mes_normalizado in mapa_meses:
            mes_atual = mapa_meses[mes_normalizado]
            continue

        if mes_atual is None:
            continue

        # --------------------------------------------------------
        # IDENTIFICAR RUBRICA
        # --------------------------------------------------------

        correspondencia = padrao_rubrica.match(linha)

        if not correspondencia:
            continue

        codigo = correspondencia.group(1)
        descricao = correspondencia.group(2).strip()
        valor = normalizar_valor(
            correspondencia.group(3)
        )

        # --------------------------------------------------------
        # ESCOLHER VENCIMENTO OU DESCONTO
        # --------------------------------------------------------

        if codigo in codigos_desconto:
            destino = rubricas_descontos
        else:
            destino = rubricas_vencimentos

        chave = (codigo, descricao)

        if chave not in destino:

            destino[chave] = {
                "Código": codigo,
                "Descrição": descricao,
                **{
                    coluna: 0.0
                    for coluna in COLUNAS_VALORES
                }
            }

        destino[chave][mes_atual] = valor

    return (
        list(rubricas_vencimentos.values()),
        list(rubricas_descontos.values()),
    )


# ============================================================
# LOCALIZAR FIM DO TEMPLATE
# ============================================================

def encontrar_ultima_linha_preenchida(ws) -> int:
    """
    Encontra a última linha que realmente possui conteúdo
    no template.

    Isso evita depender de uma linha fixa como 33.
    """

    ultima_linha = 0

    for row in ws.iter_rows():

        for celula in row:

            if celula.value is not None:
                ultima_linha = max(
                    ultima_linha,
                    celula.row
                )

    return ultima_linha


# ============================================================
# ESCREVER TABELA
# ============================================================

def escrever_tabela(
    ws,
    linha_inicial: int,
    titulo: str,
    rubricas: list[dict],
) -> int:
    """Escreve uma tabela de rubricas e retorna a próxima linha."""

    # --------------------------------------------------------
    # TÍTULO
    # --------------------------------------------------------

    ws.cell(
        linha_inicial,
        1,
        titulo
    )

    ws.cell(
        linha_inicial,
        1
    ).font = Font(
        bold=True,
        size=13
    )

    linha_tabela = linha_inicial + 1

    # --------------------------------------------------------
    # CABEÇALHO
    # --------------------------------------------------------

    cabecalho = [
        "Código",
        "Descrição",
        *COLUNAS_VALORES,
    ]

    for coluna, valor in enumerate(
        cabecalho,
        start=1
    ):

        celula = ws.cell(
            linha_tabela,
            coluna,
            valor
        )

        celula.font = Font(
            bold=True,
            color="FFFFFF"
        )

        celula.fill = PatternFill(
            "solid",
            fgColor="EB7B71"
        )

        celula.alignment = Alignment(
            horizontal="center"
        )

    # --------------------------------------------------------
    # RUBRICAS
    # --------------------------------------------------------

    for i, rubrica in enumerate(
        rubricas,
        start=linha_tabela + 1
    ):

        for coluna, campo in enumerate(
            cabecalho,
            start=1
        ):

            ws.cell(
                i,
                coluna,
                rubrica[campo]
            )

    # --------------------------------------------------------
    # FORMATAÇÃO FINANCEIRA
    # --------------------------------------------------------

    primeira_coluna_valor = 3

    ultima_linha = (
        linha_tabela + len(rubricas)
    )

    for row in ws.iter_rows(
        min_row=linha_tabela + 1,
        max_row=ultima_linha,
        min_col=primeira_coluna_valor,
        max_col=len(cabecalho),
    ):
        for celula in row:
            celula.number_format = '#,##0.00'

    # ========================================================
    # LINHA DE TOTAL
    # ========================================================

    linha_total = ultima_linha + 1

    ws.cell(
        linha_total,
        1,
        f"TOTAL {titulo}"
    )

    ws.cell(
        linha_total,
        1
    ).font = Font(
        bold=True
    )

    for coluna in range(
        primeira_coluna_valor,
        len(cabecalho) + 1
    ):

        letra_coluna = get_column_letter(
            coluna
        )

        ws.cell(
            linha_total,
            coluna,
            f"=SUM({letra_coluna}{linha_tabela + 1}:{letra_coluna}{ultima_linha})"
        )

        ws.cell(
            linha_total,
            coluna
        ).number_format = '#,##0.00'

        ws.cell(
            linha_total,
            coluna
        ).font = Font(
            bold=True
        )

    return linha_total + 2


def separar_fichas(texto: str) -> list[str]:
    """
    Separa fichas financeiras individuais.

    Suporta:
    - Modelo novo: Ficha Financeira Resumo Geral do Ano de YYYY
    - Modelo antigo: Relatório da Ficha Financeira YYYY
    """

    # ============================================================
    # MODELO ANTIGO
    # ============================================================

    if re.search(
        r"Relatório\s+da\s+Ficha\s+Financeira\s+\d{4}",
        texto,
        flags=re.IGNORECASE
    ):
        return [texto.strip()]

    # ============================================================
    # MODELO NOVO
    # ============================================================

    padrao = re.compile(
        r"Ficha\s+Financeira\s+Resumo\s+Geral\s+do\s+Ano\s+de\s+\d{4}",
        flags=re.IGNORECASE
    )

    ocorrencias = list(
        padrao.finditer(texto)
    )

    if not ocorrencias:
        raise ValueError(
            "Não foi encontrada nenhuma ficha financeira no PDF."
        )

    fichas = []

    for i, ocorrencia in enumerate(ocorrencias):

        inicio = ocorrencia.start()

        if i + 1 < len(ocorrencias):
            fim = ocorrencias[i + 1].start()
        else:
            fim = len(texto)

        trecho = texto[inicio:fim].strip()

        possui_vencimentos = re.search(
            r"Total\s+de\s+Vencimentos",
            trecho,
            flags=re.IGNORECASE
        )

        possui_descontos = re.search(
            r"Total\s+de\s+Descontos",
            trecho,
            flags=re.IGNORECASE
        )

        if possui_vencimentos and possui_descontos:
            fichas.append(trecho)

    if not fichas:
        raise ValueError(
            "Não foi encontrada nenhuma ficha financeira válida."
        )

    return fichas


# ============================================================
# EXTRAIR ANO DO PDF
# ============================================================

def extrair_ano_pdf(texto: str) -> int:
    """
    Extrai o ano da ficha financeira.

    Aceita diferentes formatos de cabeçalho:
    - Ficha Financeira Resumo Geral do Ano de 2014
    - Relatório da Ficha Financeira 2002
    """

    padroes = [
        r"Ficha\s+Financeira\s+Resumo\s+Geral\s+do\s+Ano\s+de\s+(\d{4})",
        r"Relatório\s+da\s+Ficha\s+Financeira\s+(\d{4})",
    ]

    for padrao in padroes:
        resultado = re.search(
            padrao,
            texto,
            flags=re.IGNORECASE
        )

        if resultado:
            return int(resultado.group(1))

    raise ValueError(
        "Não foi possível identificar o ano da ficha financeira."
    )

# ============================================================
# PROCESSAR PASTA
# ============================================================

def processar_pasta(
    caminho_pasta: Path,
    caminho_template: Path,
) -> None:

    print(
        f"Pasta de entrada: {caminho_pasta}"
    )

    # --------------------------------------------------------
    # LOCALIZAR PDFs
    # --------------------------------------------------------

    arquivos_pdf = sorted(
        caminho_pasta.glob("*.pdf")
    )

    if not arquivos_pdf:
        raise ValueError(
            "Nenhum arquivo PDF foi encontrado na pasta."
        )

    print(
        f"PDFs encontrados: {len(arquivos_pdf)}"
    )

    # --------------------------------------------------------
    # LER TODOS OS PDFs
    # --------------------------------------------------------

    textos_pdfs = []

    for caminho_pdf in arquivos_pdf:

        print(
            f"Lendo: {caminho_pdf.name}"
        )

        texto = extrair_texto_pdf(
            caminho_pdf
        )

        ano = extrair_ano_pdf(
            texto
        )

        textos_pdfs.append(
            {
                "arquivo": caminho_pdf,
                "texto": texto,
                "ano": ano,
            }
        )

    # --------------------------------------------------------
    # ABRIR TEMPLATE
    # --------------------------------------------------------

    wb = load_workbook(
        caminho_template,
        keep_vba=True,
        data_only=False
    )

    if NOME_ABA not in wb.sheetnames:
        raise ValueError(
            f"A aba '{NOME_ABA}' não foi encontrada no template."
        )

    # --------------------------------------------------------
    # ABA MODELO
    # --------------------------------------------------------

    ws_template = wb[NOME_ABA]

    # --------------------------------------------------------
    # DESCOBRIR OS ANOS
    # --------------------------------------------------------

    anos = []

    for item in textos_pdfs:

        ano = item["ano"]

        if ano not in anos:
            anos.append(ano)

    anos.sort()

    print(
        f"\nAnos encontrados: {', '.join(map(str, anos))}"
    )

    # --------------------------------------------------------
    # CRIAR TODAS AS ABAS ANTES DE PREENCHER
    # --------------------------------------------------------

    abas_por_ano = {}

    for indice, ano in enumerate(anos):

        nome_aba = str(ano)

        if indice == 0:

            # A primeira aba utiliza o próprio template.
            ws = ws_template

            ws.title = nome_aba

        else:

            # As demais abas são cópias do template
            # ainda vazio.
            ws = wb.copy_worksheet(
                ws_template
            )

            ws.title = nome_aba

        abas_por_ano[nome_aba] = ws

    # --------------------------------------------------------
    # PROCESSAR OS PDFs
    # --------------------------------------------------------

    for numero_pdf, item in enumerate(
        textos_pdfs,
        start=1
    ):

        caminho_pdf = item["arquivo"]
        texto = item["texto"]
        ano = item["ano"]

        print(
            f"\n{'=' * 60}"
        )

        print(
            f"PDF {numero_pdf}/{len(textos_pdfs)}"
        )

        print(
            f"Arquivo: {caminho_pdf.name}"
        )

        print(
            f"Ano: {ano}"
        )

        # ----------------------------------------------------
        # PLANILHA DO ANO
        # ----------------------------------------------------

        ws = abas_por_ano[
            str(ano)
        ]

        # ----------------------------------------------------
        # ENCONTRAR PRÓXIMA LINHA
        # ----------------------------------------------------

        ultima_linha = encontrar_ultima_linha_preenchida(
            ws
        )

        linha = ultima_linha + 5

        # ----------------------------------------------------
        # SEPARAR FICHAS
        # ----------------------------------------------------

        fichas = separar_fichas(
            texto
        )

        print(
            f"Fichas encontradas: {len(fichas)}"
        )

        # ----------------------------------------------------
        # PROCESSAR CADA FICHA
        # ----------------------------------------------------

        for numero_ficha, texto_ficha in enumerate(
            fichas,
            start=1
        ):

            print(
                f"\nProcessando ficha {numero_ficha}..."
            )

            # ------------------------------------------------
            # DADOS DO SERVIDOR
            # ------------------------------------------------

            if re.search(
                r"Relatório\s+da\s+Ficha\s+Financeira\s+\d{4}",
                texto_ficha,
                flags=re.IGNORECASE
            ):
                dados_servidor = extrair_dados_servidor_modelo_antigo(
                    texto_ficha
                )
            else:
                dados_servidor = extrair_dados_servidor(
                    texto_ficha
                )

            # ------------------------------------------------
            # RUBRICAS
            # ------------------------------------------------

            if re.search(
                r"Relatório\s+da\s+Ficha\s+Financeira\s+\d{4}",
                texto_ficha,
                flags=re.IGNORECASE
            ):
                (
                    rubricas_vencimentos,
                    rubricas_descontos,
                ) = extrair_rubricas_modelo_antigo(
                    texto_ficha
                )
            else:
                (
                    rubricas_vencimentos,
                    rubricas_descontos,
                ) = extrair_rubricas(
                    texto_ficha
                )

            if not dados_servidor:

                raise ValueError(
                    f"Não foi possível identificar os dados "
                    f"do servidor da ficha {numero_ficha}."
                )

            if (
                not rubricas_vencimentos
                and not rubricas_descontos
            ):

                raise ValueError(
                    f"Nenhuma rubrica foi encontrada "
                    f"na ficha {numero_ficha}."
                )

            # ------------------------------------------------
            # DADOS DO SERVIDOR
            # ------------------------------------------------

            coluna_dados = (
                len(
                    [
                        "Código",
                        "Descrição",
                        *COLUNAS_VALORES
                    ]
                ) + 2
            )

            ws.cell(
                linha,
                coluna_dados,
                f"DADOS DO SERVIDOR - FICHA {numero_ficha}"
            )

            ws.cell(
                linha,
                coluna_dados
            ).font = Font(
                bold=True,
                size=14
            )

            linha_dados = linha + 2

            for campo, valor in dados_servidor.items():

                ws.cell(
                    linha_dados,
                    coluna_dados,
                    campo
                ).font = Font(
                    bold=True
                )

                ws.cell(
                    linha_dados,
                    coluna_dados + 1,
                    valor
                )

                linha_dados += 1

            # ------------------------------------------------
            # VENCIMENTOS
            # ------------------------------------------------

            linha = escrever_tabela(
                ws,
                linha,
                f"VENCIMENTOS - FICHA {numero_ficha}",
                rubricas_vencimentos,
            )

            # ------------------------------------------------
            # DESCONTOS
            # ------------------------------------------------

            linha = escrever_tabela(
                ws,
                linha,
                f"DESCONTOS - FICHA {numero_ficha}",
                rubricas_descontos,
            )

            # ------------------------------------------------
            # INFORMAÇÕES
            # ------------------------------------------------

            print(
                f"  Servidor: "
                f"{dados_servidor.get('Nome', 'não identificado')}"
            )

            print(
                f"  Cargo: "
                f"{dados_servidor.get('Cargo', 'não identificado')}"
            )

            print(
                f"  Vencimentos: "
                f"{len(rubricas_vencimentos)} rubricas"
            )

            print(
                f"  Descontos: "
                f"{len(rubricas_descontos)} rubricas"
            )

    # --------------------------------------------------------
    # REMOVER ABA MODELO
    # --------------------------------------------------------

    # Depois que todas as abas foram criadas,
    # Plan1 já não é mais necessária.
    #
    # Porém, se o primeiro ano utilizou a própria Plan1,
    # ela já foi renomeada e não existe mais como Plan1.
    if NOME_ABA in wb.sheetnames:

        del wb[NOME_ABA]

    # --------------------------------------------------------
    # RECÁLCULO DAS FÓRMULAS
    # --------------------------------------------------------

    try:

        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"

    except AttributeError:

        pass

    # --------------------------------------------------------
    # SALVAR
    # --------------------------------------------------------

    caminho_saida = (
        caminho_pasta / "Ficha Financeira.xlsm"
    )

    wb.save(
        caminho_saida
    )

    print(
        f"\n{'=' * 60}"
    )

    print(
        f"Excel criado: {caminho_saida}"
    )

    print(
        f"Abas criadas: {', '.join(abas_por_ano.keys())}"
    )

# ============================================================
# MAIN
# ============================================================

def main():

    if len(sys.argv) != 2:
        print(
            "Uso: python src/main.py pasta"
        )
        sys.exit(1)

    caminho_entrada = Path(
        sys.argv[1]
    )

    # --------------------------------------------------------
    # VERIFICAR ENTRADA
    # --------------------------------------------------------

    if not caminho_entrada.exists():
        print(
            f"Arquivo ou pasta não encontrado: "
            f"{caminho_entrada}"
        )
        sys.exit(1)

    if not caminho_entrada.is_dir():
        print(
            "A entrada informada não é uma pasta."
        )
        sys.exit(1)

    # --------------------------------------------------------
    # TEMPLATE
    # --------------------------------------------------------

    if not ARQUIVO_TEMPLATE.exists():
        print(
            f"Template não encontrado: "
            f"{ARQUIVO_TEMPLATE}"
        )

        print(
            "Coloque o AnexoVMacro.xlsm na pasta "
            "data."
        )

        sys.exit(1)

    # --------------------------------------------------------
    # PROCESSAR
    # --------------------------------------------------------

    try:

        processar_pasta(
            caminho_entrada,
            ARQUIVO_TEMPLATE,
        )

    except Exception as erro:

        print(
            f"\nErro: {erro}"
        )

        sys.exit(1)


if __name__ == "__main__":
    main()